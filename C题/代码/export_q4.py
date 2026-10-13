"""Export and validate the fluctuating-price Problem 4 workbooks."""

from __future__ import annotations

import argparse
import json
from copy import copy
from datetime import datetime
from pathlib import Path

import numpy as np
from openpyxl import load_workbook


PERIODS = 144
BLOCKS = 6
BLOCK_SIZE = 24
REPORTING_TOLERANCE = 1e-3
BATTERY_MIN = 1200.0
BATTERY_MAX = 10800.0
BATTERY_POWER = 5000.0 / 6.0
EMERGENCY_MULTIPLIER = 5.0
LABELS_4H = [
    "0:00-4:00", "4:00-8:00", "8:00-12:00",
    "12:00-16:00", "16:00-20:00", "20:00-24:00",
]


def _read(path: Path) -> dict:
    return json.loads(path.read_text(encoding="utf-8"))


def _normalize_label(value):
    if value == "7:0-7:10":
        return "7:00-7:10"
    if value == "0:00-0:10+1":
        return "0:00+1-0:10+1"
    return value


def _header(ws):
    return {_normalize_label(ws.cell(1, c).value): c for c in range(1, ws.max_column + 1)}


def _sum_block(values, block):
    start = block * BLOCK_SIZE
    return float(sum(values[start:start + BLOCK_SIZE]))


def _interval_rows(days):
    rows = []
    for day in days:
        intervals = day.get("emergency_intervals", [])
        if not intervals:
            rows.append((datetime.fromisoformat(day["date"]), None, None, 0, 1))
            continue
        for i, item in enumerate(intervals):
            rows.append((
                datetime.fromisoformat(day["date"]) if i == 0 else None,
                item["time_period"],
                float(item["purchase_kwh"]),
                i,
                len(intervals),
            ))
    return rows


def _clear_rows(ws, first_row, last_row, columns):
    for row in range(first_row, last_row + 1):
        for col in range(1, columns + 1):
            ws.cell(row, col).value = None


def _write_plan(ws, solution, value_key, cost_key):
    days = solution["days"]
    labels = solution["interval_labels"]
    header = _header(ws)
    missing = [label for label in labels if label not in header]
    if missing:
        raise ValueError(f"模板缺少时间标签: {missing[:3]}")
    if len(days) != 334 or ws.max_row < 335 or ws.max_column < 147:
        raise ValueError("问题四计划表尺寸或天数不正确")
    for row, day in enumerate(days, start=2):
        values = day[value_key]
        if len(values) != PERIODS:
            raise ValueError(f"{day['date']}的{value_key}长度不为144")
        ws.cell(row, 1).value = datetime.fromisoformat(day["date"])
        ws.cell(row, 1).number_format = "yyyy/m/d"
        for label, value in zip(labels, values):
            ws.cell(row, header[label]).value = float(value)
            ws.cell(row, header[label]).number_format = "0.0000"
        ws.cell(row, 146).value = float(sum(values))
        ws.cell(row, 146).number_format = "0.0000"
        ws.cell(row, 147).value = float(day[cost_key])
        ws.cell(row, 147).number_format = "0.0000"


def _write_battery(ws, days, charge_key="charge_kwh", discharge_key="discharge_kwh", soc_key="soc_kwh"):
    styles = [[copy(ws.cell(r, c)._style) for c in range(1, 7)] for r in range(2, 8)]
    _clear_rows(ws, 2, ws.max_row, 6)
    required = 1 + len(days) * BLOCKS
    if ws.max_row < required:
        ws.insert_rows(ws.max_row + 1, required - ws.max_row)
    for day_index, day in enumerate(days):
        charge = day[charge_key]
        discharge = day[discharge_key]
        soc = day[soc_key]
        if len(charge) != PERIODS or len(discharge) != PERIODS or len(soc) != PERIODS + 1:
            raise ValueError(f"{day['date']}的储能数组长度错误")
        for block in range(BLOCKS):
            row = 2 + day_index * BLOCKS + block
            for col in range(1, 7):
                ws.cell(row, col)._style = copy(styles[block][col - 1])
            ws.cell(row, 1).value = datetime.fromisoformat(day["date"]) if block == 0 else None
            ws.cell(row, 1).number_format = "yyyy/m/d"
            ws.cell(row, 2).value = LABELS_4H[block]
            ws.cell(row, 3).value = _sum_block(charge, block)
            ws.cell(row, 4).value = _sum_block(discharge, block)
            ws.cell(row, 3).number_format = ws.cell(row, 4).number_format = "0.0000"
            ws.cell(row, 5).value = "0:00" if block == 0 else "24:00" if block == 1 else None
            ws.cell(row, 6).value = float(soc[0]) if block == 0 else float(soc[-1]) if block == 1 else None
            ws.cell(row, 6).number_format = "0.0000"
    if ws.max_row > required:
        ws.delete_rows(required + 1, ws.max_row - required)


def _write_emergency(ws, days):
    styles = [[copy(ws.cell(r, c)._style) for c in range(1, 4)] for r in range(2, 5)]
    rows = _interval_rows(days)
    _clear_rows(ws, 2, ws.max_row, 3)
    required = 1 + len(rows)
    if ws.max_row < required:
        ws.insert_rows(ws.max_row + 1, required - ws.max_row)
    for i, (dt, period, amount, group_i, group_n) in enumerate(rows, start=2):
        style_i = 0 if group_i == 0 else 2 if group_i == group_n - 1 else 1
        for col in range(1, 4):
            ws.cell(i, col)._style = copy(styles[style_i][col - 1])
        ws.cell(i, 1).value = dt
        ws.cell(i, 1).number_format = "yyyy/m/d"
        ws.cell(i, 2).value = period
        ws.cell(i, 3).value = amount
        ws.cell(i, 3).number_format = "0.0000"
    if ws.max_row > required:
        ws.delete_rows(required + 1, ws.max_row - required)


def _verify_plan_sheet(ws, days, labels, value_key, cost_key):
    """Read a 144-slot purchase sheet back and compare every written cell."""
    if ws.max_row != len(days) + 1 or ws.max_column < 147:
        raise AssertionError(f"{ws.title}计划表尺寸错误")
    header = _header(ws)
    missing = [label for label in labels if label not in header]
    if missing:
        raise AssertionError(f"{ws.title}缺少时间标签: {missing[:3]}")
    for day_index, day in enumerate(days):
        values = np.asarray(day[value_key], dtype=float)
        if values.shape != (PERIODS,):
            raise AssertionError(f"{day['date']}的{value_key}长度错误")
        row = day_index + 2
        cell_date = ws.cell(row, 1).value
        if cell_date is None or not hasattr(cell_date, "date"):
            raise AssertionError(f"{ws.title}!A{row}日期为空")
        if cell_date.date().isoformat() != day["date"]:
            raise AssertionError(f"{ws.title}第{row}行日期写入错误")
        for label, value in zip(labels, values):
            cell = ws.cell(row, header[label]).value
            if cell is None or abs(float(cell) - float(value)) > 1e-9:
                raise AssertionError(f"{ws.title}{day['date']} {label}数值写入错误")
        total = ws.cell(row, 146).value
        if total is None or abs(float(total) - float(values.sum())) > 1e-9:
            raise AssertionError(f"{ws.title}{day['date']}购电合计写入错误")
        cost = ws.cell(row, 147).value
        if cost is None or abs(float(cost) - float(day[cost_key])) > 1e-9:
            raise AssertionError(f"{ws.title}{day['date']}购电费用写入错误")


def _verify_battery_sheet(ws, days):
    """Read back the six four-hour battery rows written for each day."""
    required = 1 + len(days) * BLOCKS
    if ws.max_row != required or ws.max_column < 6:
        raise AssertionError(f"{ws.title}储能表尺寸错误")
    for day_index, day in enumerate(days):
        charge = np.asarray(day["charge_kwh"], dtype=float)
        discharge = np.asarray(day["discharge_kwh"], dtype=float)
        soc = np.asarray(day["soc_kwh"], dtype=float)
        for block in range(BLOCKS):
            row = 2 + day_index * BLOCKS + block
            expected_label = LABELS_4H[block]
            if ws.cell(row, 2).value != expected_label:
                raise AssertionError(f"{ws.title}!B{row}时间块写入错误")
            date_value = ws.cell(row, 1).value
            if block == 0:
                if date_value is None or not hasattr(date_value, "date"):
                    raise AssertionError(f"{ws.title}!A{row}日期为空")
                if date_value.date().isoformat() != day["date"]:
                    raise AssertionError(f"{ws.title}!A{row}日期写入错误")
            elif date_value is not None:
                raise AssertionError(f"{ws.title}!A{row}非首块不应写日期")
            expected_charge = _sum_block(charge, block)
            expected_discharge = _sum_block(discharge, block)
            actual_charge = ws.cell(row, 3).value
            actual_discharge = ws.cell(row, 4).value
            if actual_charge is None or abs(float(actual_charge) - expected_charge) > 1e-9:
                raise AssertionError(f"{ws.title}{day['date']} {expected_label}充电合计错误")
            if actual_discharge is None or abs(float(actual_discharge) - expected_discharge) > 1e-9:
                raise AssertionError(f"{ws.title}{day['date']} {expected_label}放电合计错误")
            expected_time = "0:00" if block == 0 else "24:00" if block == 1 else None
            if ws.cell(row, 5).value != expected_time:
                raise AssertionError(f"{ws.title}{day['date']} {expected_label}时间边界错误")
            expected_soc = float(soc[0]) if block == 0 else float(soc[-1]) if block == 1 else None
            actual_soc = ws.cell(row, 6).value
            if expected_soc is None:
                if actual_soc is not None:
                    raise AssertionError(f"{ws.title}{day['date']} {expected_label}SOC边界应为空")
            elif actual_soc is None or abs(float(actual_soc) - expected_soc) > 1e-9:
                raise AssertionError(f"{ws.title}{day['date']} {expected_label}SOC边界错误")


def _verify_emergency_sheet(ws, days):
    """Read back every displayed emergency interval, including its date."""
    expected_rows = _interval_rows(days)
    if ws.max_row != len(expected_rows) + 1 or ws.max_column < 3:
        raise AssertionError("紧急购电表尺寸错误")
    for row, expected in enumerate(expected_rows, start=2):
        actual_date, actual_period, actual_amount = (
            ws.cell(row, 1).value,
            ws.cell(row, 2).value,
            ws.cell(row, 3).value,
        )
        expected_date, expected_period, expected_amount = expected[:3]
        if expected_date is None:
            if actual_date is not None:
                raise AssertionError(f"紧急购电表第{row}行日期分组错误")
        elif actual_date is None or not hasattr(actual_date, "date"):
            raise AssertionError(f"紧急购电表第{row}行日期为空")
        elif actual_date.date() != expected_date.date():
            raise AssertionError(f"紧急购电表第{row}行日期错误")
        if actual_period != expected_period:
            raise AssertionError(f"紧急购电表第{row}行时间段错误")
        if expected_amount is None:
            if actual_amount not in (None, 0, 0.0):
                raise AssertionError(f"紧急购电表第{row}行空记录数值错误")
        elif actual_amount is None or abs(float(actual_amount) - float(expected_amount)) > 1e-9:
            raise AssertionError(f"紧急购电表第{row}行购电量错误")


def export_problem42(c_root: Path) -> Path:
    solution = _read(c_root / "结果" / "q4_2_solution.json")
    wb = load_workbook(c_root / "附件" / "附件5" / "result4-2.xlsx")
    plan, battery, emergency = wb.worksheets[:3]
    _write_plan(plan, solution, "planned_purchase_kwh", "planned_cost_yuan")
    _write_battery(battery, solution["days"])
    _write_emergency(emergency, solution["days"])
    output = c_root / "结果" / "result4-2.xlsx"
    wb.save(output)
    wb.close()
    return output


def export_problem43(c_root: Path) -> Path:
    solution = _read(c_root / "结果" / "q4_3_solution.json")
    wb = load_workbook(c_root / "附件" / "附件5" / "result4-3.xlsx")
    plan, adjusted, battery, emergency = wb.worksheets[:4]
    _write_plan(plan, solution, "planned_purchase_kwh", "planned_cost_yuan")
    _write_plan(adjusted, solution, "adjusted_purchase_kwh", "normal_settlement_cost_yuan")
    _write_battery(battery, solution["days"])
    _write_emergency(emergency, solution["days"])
    output = c_root / "结果" / "result4-3.xlsx"
    wb.save(output)
    wb.close()
    return output


def _verify_common(solution, workbook, has_adjustment):
    days = solution["days"]
    if len(days) != 334:
        raise AssertionError("问题四结果必须包含334天")
    labels = solution["interval_labels"]
    if len(labels) != PERIODS:
        raise AssertionError("问题四时间标签数量错误")
    plan = workbook.worksheets[0]
    _verify_plan_sheet(plan, days, labels, "planned_purchase_kwh", "planned_cost_yuan")
    battery = workbook.worksheets[2 if has_adjustment else 1]
    _verify_battery_sheet(battery, days)
    if has_adjustment:
        adjusted_sheet = workbook.worksheets[1]
        _verify_plan_sheet(
            adjusted_sheet,
            days,
            labels,
            "adjusted_purchase_kwh",
            "normal_settlement_cost_yuan",
        )
    emergency_sheet = workbook.worksheets[3 if has_adjustment else 2]
    _verify_emergency_sheet(emergency_sheet, days)
    labels = solution["interval_labels"]
    for day_index, day in enumerate(days):
        price = np.asarray(day["price_yuan_per_kwh"], dtype=float)
        planned = np.asarray(day["planned_purchase_kwh"], dtype=float)
        charge = np.asarray(day["charge_kwh"], dtype=float)
        discharge = np.asarray(day["discharge_kwh"], dtype=float)
        emergency = np.asarray(day["emergency_kwh"], dtype=float)
        soc = np.asarray(day["soc_kwh"], dtype=float)
        if any(x.shape != (PERIODS,) for x in (price, planned, charge, discharge, emergency)) or soc.shape != (PERIODS + 1,):
            raise AssertionError(f"{day['date']}逐时数组长度错误")
        if float(day["max_balance_error_kwh"]) > 1e-7:
            raise AssertionError(f"{day['date']}供需平衡误差超限")
        if soc.min() < BATTERY_MIN - 1e-7 or soc.max() > BATTERY_MAX + 1e-7:
            raise AssertionError(f"{day['date']}储能越界")
        if np.max(charge + discharge) > BATTERY_POWER + 1e-7:
            raise AssertionError(f"{day['date']}充放电功率越界")
        if np.any((charge > 1e-7) & (discharge > 1e-7)):
            raise AssertionError(f"{day['date']}存在同时充放电")
        if abs(float(day["planned_cost_yuan"]) - float(np.dot(price, planned))) > 1e-6:
            raise AssertionError(f"{day['date']}计划购电费与波动电价不一致")
        if abs(float(day["emergency_cost_yuan"]) - float(np.dot(EMERGENCY_MULTIPLIER * price, emergency))) > 1e-6:
            raise AssertionError(f"{day['date']}紧急购电费与波动电价不一致")
        if has_adjustment:
            adjusted = np.asarray(day["adjusted_purchase_kwh"], dtype=float)
            if adjusted.shape != (PERIODS,) or not np.isfinite(adjusted).all() or np.any(adjusted < -1e-9):
                raise AssertionError(f"{day['date']}调整购电数组非法")
            if abs(float(day["normal_settlement_cost_yuan"]) - (
                float(np.dot(price, np.minimum(planned, adjusted)))
                + float(np.dot(0.5 * price, np.maximum(planned - adjusted, 0.0)))
                + float(np.dot(1.5 * price, np.maximum(adjusted - planned, 0.0)))
            )) > 1e-6:
                raise AssertionError(f"{day['date']}调整结算费用与波动电价不一致")
        if day_index and abs(float(soc[0]) - float(days[day_index - 1]["soc_kwh"][-1])) > 1e-7:
            raise AssertionError(f"{day['date']}储能未连续传递")
    total = sum(float(emergency_sheet.cell(r, 3).value or 0.0) for r in range(2, emergency_sheet.max_row + 1))
    # The workbook intentionally suppresses sub-1e-3 kWh fragments when it
    # builds contiguous reporting intervals.  Validate the displayed table
    # against the same reported intervals; the JSON summary retains raw
    # machine-precision energy for independent numerical audits.
    reported_total = sum(
        float(item["purchase_kwh"])
        for day in days
        for item in day.get("emergency_intervals", [])
    )
    if abs(total - reported_total) > 1e-6:
        raise AssertionError("紧急购电表内区间合计错误")
    raw_total = float(solution["summary"]["emergency_energy_kwh"])
    if reported_total - raw_total > 1e-6:
        raise AssertionError("紧急购电表合计超过JSON原始合计")
    for ws in workbook.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and (cell.value.startswith("#") or cell.value == "⁝"):
                    raise AssertionError(f"{ws.title}!{cell.coordinate}存在错误或模板省略符")


def verify_problem42(c_root: Path) -> None:
    solution = _read(c_root / "结果" / "q4_2_solution.json")
    wb = load_workbook(c_root / "结果" / "result4-2.xlsx", data_only=False)
    _verify_common(solution, wb, False)
    wb.close()


def verify_problem43(c_root: Path) -> None:
    solution = _read(c_root / "结果" / "q4_3_solution.json")
    wb = load_workbook(c_root / "结果" / "result4-3.xlsx", data_only=False)
    _verify_common(solution, wb, True)
    wb.close()


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("question", choices=("2", "3", "all"), nargs="?", default="all")
    args = parser.parse_args()
    root = Path(__file__).resolve().parents[1]
    outputs = []
    if args.question in ("2", "all"):
        outputs.append(export_problem42(root))
        verify_problem42(root)
    if args.question in ("3", "all"):
        outputs.append(export_problem43(root))
        verify_problem43(root)
    for output in outputs:
        print(f"已生成并校验：{output}")


if __name__ == "__main__":
    main()
