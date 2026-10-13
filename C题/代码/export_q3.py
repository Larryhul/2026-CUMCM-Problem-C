"""Export Question 3 JSON into the supplied result3.xlsx template."""

from __future__ import annotations

import argparse
import json
import shutil
from copy import copy
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

PERIODS = 144
BLOCKS = 6
BLOCK_SIZE = 24
LABELS = [
    "0:00-4:00", "4:00-8:00", "8:00-12:00", "12:00-16:00",
    "16:00-20:00", "20:00-24:00",
]
TOL = 1e-3


def _sum_block(values, block):
    return float(sum(values[block * BLOCK_SIZE:(block + 1) * BLOCK_SIZE]))


def _interval_rows(days):
    rows = []
    for day in days:
        intervals = day.get("emergency_intervals", [])
        if not intervals:
            rows.append((datetime.fromisoformat(day["date"]), None, None, 0, 1))
        else:
            for i, interval in enumerate(intervals):
                rows.append((
                    datetime.fromisoformat(day["date"]) if i == 0 else None,
                    interval["time_period"], float(interval["purchase_kwh"]), i, len(intervals),
                ))
    return rows


def _copy_style(ws, src_row, dst_row, max_col):
    for col in range(1, max_col + 1):
        ws.cell(dst_row, col)._style = copy(ws.cell(src_row, col)._style)
        if ws.cell(src_row, col).has_style:
            ws.cell(dst_row, col).number_format = ws.cell(src_row, col).number_format


def export(c_root: Path) -> Path:
    solution = json.loads((c_root / "结果" / "q3_solution.json").read_text(encoding="utf-8"))
    days = solution["days"]
    if len(days) != 334:
        raise ValueError(f"问题三应包含334天，实际{len(days)}天")
    template = c_root / "附件" / "附件5" / "result3.xlsx"
    wb = load_workbook(template)
    plan_ws = wb["计划购电量"]
    adj_ws = wb["调整购电量"]
    batt_ws = wb["充放电量"]
    em_ws = wb["紧急购电量"]

    # The first two sheets have a stable 334-day layout and 144 interval
    # columns followed by daily totals.
    def normalize_label(value):
        if value == "7:0-7:10":
            return "7:00-7:10"
        if value == "0:00-0:10+1":
            return "0:00+1-0:10+1"
        return value
    header = {normalize_label(plan_ws.cell(1, c).value): c for c in range(1, plan_ws.max_column + 1)}
    labels = solution["interval_labels"]
    missing = [x for x in labels if x not in header]
    if missing:
        raise ValueError(f"result3.xlsx缺少时间标签: {missing[:3]}")
    for row, day in enumerate(days, start=2):
        dt = datetime.fromisoformat(day["date"])
        for ws, key, cost_key in ((plan_ws, "planned_purchase_kwh", "planned_cost_yuan"),
                                  (adj_ws, "adjusted_purchase_kwh", "normal_settlement_cost_yuan")):
            ws.cell(row, 1).value = dt
            ws.cell(row, 1).number_format = "yyyy/m/d"
            values = day[key]
            if len(values) != PERIODS:
                raise ValueError(f"{day['date']} {key}长度错误")
            for label, value in zip(labels, values):
                ws.cell(row, header[label]).value = float(value)
                ws.cell(row, header[label]).number_format = "0.0000"
            ws.cell(row, 146).value = float(sum(values))
            ws.cell(row, 146).number_format = "0.0000"
            ws.cell(row, 147).value = float(day[cost_key])
            ws.cell(row, 147).number_format = "0.0000"

    # Expand the four-hour storage table using the six template style rows.
    battery_styles = [[copy(batt_ws.cell(r, c)._style) for c in range(1, 7)] for r in range(2, 8)]
    for r in range(2, batt_ws.max_row + 1):
        for c in range(1, 7):
            batt_ws.cell(r, c).value = None
    required = 1 + len(days) * BLOCKS
    if batt_ws.max_row < required:
        batt_ws.insert_rows(batt_ws.max_row + 1, required - batt_ws.max_row)
    for di, day in enumerate(days):
        row0 = 2 + di * BLOCKS
        charges, discharges, soc = day["charge_kwh"], day["discharge_kwh"], day["soc_kwh"]
        for b in range(BLOCKS):
            row = row0 + b
            for c in range(1, 7):
                batt_ws.cell(row, c)._style = copy(battery_styles[b][c - 1])
            batt_ws.cell(row, 1).value = datetime.fromisoformat(day["date"]) if b == 0 else None
            batt_ws.cell(row, 1).number_format = "yyyy/m/d"
            batt_ws.cell(row, 2).value = LABELS[b]
            batt_ws.cell(row, 3).value = _sum_block(charges, b)
            batt_ws.cell(row, 4).value = _sum_block(discharges, b)
            batt_ws.cell(row, 3).number_format = batt_ws.cell(row, 4).number_format = "0.0000"
            batt_ws.cell(row, 5).value = "0:00" if b == 0 else "24:00" if b == 1 else None
            batt_ws.cell(row, 6).value = float(soc[0]) if b == 0 else float(soc[-1]) if b == 1 else None
            batt_ws.cell(row, 6).number_format = "0.0000"
    if batt_ws.max_row > required:
        batt_ws.delete_rows(required + 1, batt_ws.max_row - required)

    # Expand emergency table, preserving the template's three boundary styles.
    em_styles = [[copy(em_ws.cell(r, c)._style) for c in range(1, 4)] for r in range(2, 5)]
    rows = _interval_rows(days)
    required_em = 1 + len(rows)
    for r in range(2, em_ws.max_row + 1):
        for c in range(1, 4):
            em_ws.cell(r, c).value = None
    if em_ws.max_row < required_em:
        em_ws.insert_rows(em_ws.max_row + 1, required_em - em_ws.max_row)
    for i, (dt, period, amount, group_i, group_n) in enumerate(rows, start=2):
        style_i = 0 if group_i == 0 else 2 if group_i == group_n - 1 else 1
        for c in range(1, 4):
            em_ws.cell(i, c)._style = copy(em_styles[style_i][c - 1])
        em_ws.cell(i, 1).value = dt
        em_ws.cell(i, 1).number_format = "yyyy/m/d"
        em_ws.cell(i, 2).value = period
        em_ws.cell(i, 3).value = amount
        em_ws.cell(i, 3).number_format = "0.0000"
    if em_ws.max_row > required_em:
        em_ws.delete_rows(required_em + 1, em_ws.max_row - required_em)

    out = c_root / "结果" / "result3_final.xlsx"
    official = c_root / "结果" / "result3.xlsx"
    wb.save(out)
    wb.close()
    shutil.copyfile(out, official)
    return out


def verify(c_root: Path) -> None:
    sol = json.loads((c_root / "结果" / "q3_solution.json").read_text(encoding="utf-8"))
    wb = load_workbook(c_root / "结果" / "result3_final.xlsx", data_only=False)
    assert wb["计划购电量"].max_row == 335 and wb["调整购电量"].max_row == 335
    assert wb["充放电量"].max_row == 1 + 334 * BLOCKS
    rows = _interval_rows(sol["days"])
    assert wb["紧急购电量"].max_row == len(rows) + 1
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and (cell.value.startswith("#") or cell.value == "⁝"):
                    raise AssertionError(f"{ws.title}!{cell.coordinate}存在错误或模板省略符")
    wb.close()
    official = load_workbook(c_root / "结果" / "result3.xlsx", data_only=False)
    final = load_workbook(c_root / "结果" / "result3_final.xlsx", data_only=False)
    if official.sheetnames != final.sheetnames:
        raise AssertionError("result3.xlsx与result3_final.xlsx工作表不一致")
    for name in final.sheetnames:
        ws_official = official[name]
        ws_final = final[name]
        if (ws_official.max_row, ws_official.max_column) != (ws_final.max_row, ws_final.max_column):
            raise AssertionError(f"{name}两份结果文件尺寸不一致")
        for row_official, row_final in zip(ws_official.iter_rows(), ws_final.iter_rows()):
            for cell_official, cell_final in zip(row_official, row_final):
                if cell_official.value != cell_final.value:
                    raise AssertionError(f"{name}!{cell_official.coordinate}两份结果文件内容不一致")
    official.close()
    final.close()


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--verify", action="store_true")
    args = parser.parse_args()
    root = Path(__file__).resolve().parents[1]
    out = export(root)
    if args.verify:
        verify(root)
    print(out)
