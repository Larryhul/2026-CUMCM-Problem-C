"""Write Problem 1 and Problem 2 solutions into the official Excel templates."""

from __future__ import annotations

import argparse
import json
from copy import copy
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook


PERIODS = 144
FOUR_HOUR_BLOCKS = 6
PERIODS_PER_BLOCK = 24
EXPECTED_FIRST_LABEL = "0:10-0:20"
EXPECTED_LAST_LABEL = "0:00+1-0:10+1"
FOUR_HOUR_LABELS = [
    "0:00-4:00",
    "4:00-8:00",
    "8:00-12:00",
    "12:00-16:00",
    "16:00-20:00",
    "20:00-24:00",
]


def _read_json(path: Path) -> dict:
    return json.loads(path.read_text(encoding="utf-8"))


def _validate_labels(solution: dict, name: str) -> list[str]:
    labels = solution.get("interval_labels")
    if (
        not isinstance(labels, list)
        or len(labels) != PERIODS
        or labels[0] != EXPECTED_FIRST_LABEL
        or labels[-1] != EXPECTED_LAST_LABEL
        or len(set(labels)) != PERIODS
    ):
        raise ValueError(f"{name}结果不符合统一的144时段时间规范")
    return labels


def _normalize_result2_label(value: object) -> object:
    if value == "0:00-0:10+1":
        return EXPECTED_LAST_LABEL
    if value == "7:0-7:10":
        return "7:00-7:10"
    return value


def _sum_block(values: list[float], block: int) -> float:
    start = block * PERIODS_PER_BLOCK
    return float(sum(values[start : start + PERIODS_PER_BLOCK]))


def _build_emergency_rows(
    days: list[dict],
) -> list[tuple[datetime | None, str | None, float | None, int, int]]:
    rows: list[tuple[datetime | None, str | None, float | None, int, int]] = []
    for day in days:
        intervals = day["emergency_intervals"]
        if not intervals:
            rows.append((datetime.fromisoformat(day["date"]), None, None, 0, 1))
            continue
        for index, interval in enumerate(intervals):
            rows.append(
                (
                    datetime.fromisoformat(day["date"]) if index == 0 else None,
                    interval["time_period"],
                    float(interval["purchase_kwh"]),
                    index,
                    len(intervals),
                )
            )
    return rows


def export_problem1(c_root: Path) -> Path:
    solution = _read_json(c_root / "结果" / "q1_solution.json")
    labels = _validate_labels(solution, "问题一")
    if any(len(solution[key]) != PERIODS for key in ("purchase_kwh", "charge_kwh", "discharge_kwh")):
        raise ValueError("问题一结果数组长度不为144")
    if len(solution["soc_kwh"]) != PERIODS + 1:
        raise ValueError("问题一储能状态数组长度不为145")

    workbook = load_workbook(c_root / "附件" / "附件5" / "result1.xlsx")
    plan_sheet = workbook["计划购电量"]
    battery_sheet = workbook["充放电量"]
    template_labels = [plan_sheet.cell(row, 1).value for row in range(2, PERIODS + 2)]
    if template_labels != labels:
        raise ValueError("result1.xlsx时间标签与统一规范不一致")

    for row, value in enumerate(solution["purchase_kwh"], start=2):
        plan_sheet.cell(row, 2, float(value)).number_format = "0.0000"
    for block in range(FOUR_HOUR_BLOCKS):
        row = block + 2
        battery_sheet.cell(row, 2, _sum_block(solution["charge_kwh"], block)).number_format = "0.0000"
        battery_sheet.cell(row, 3, _sum_block(solution["discharge_kwh"], block)).number_format = "0.0000"
    battery_sheet["E2"] = float(solution["soc_kwh"][0])
    battery_sheet["E3"] = float(solution["soc_kwh"][-1])
    battery_sheet["E2"].number_format = "0.0000"
    battery_sheet["E3"].number_format = "0.0000"
    battery_sheet.column_dimensions["E"].width = max(
        battery_sheet.column_dimensions["E"].width or 0,
        13,
    )

    output = c_root / "结果" / "result1.xlsx"
    workbook.save(output)
    workbook.close()
    return output


def export_problem2(c_root: Path) -> Path:
    solution = _read_json(c_root / "结果" / "q2_solution.json")
    labels = _validate_labels(solution, "问题二")
    days = solution.get("days", [])
    if len(days) != 334:
        raise ValueError(f"问题二应包含334天，实际为{len(days)}天")
    reporting_tolerance = float(
        solution.get("model", {}).get("emergency_reporting_tolerance_kwh", 0.0)
    )
    if reporting_tolerance <= 0.0:
        raise ValueError("问题二结果缺少有效的紧急购电报告阈值")

    workbook = load_workbook(c_root / "附件" / "附件5" / "result2.xlsx")
    plan_sheet = workbook["计划购电量"]
    battery_sheet = workbook["充放电量"]
    emergency_sheet = workbook["紧急购电量"]

    header_index = {
        _normalize_result2_label(plan_sheet.cell(1, column).value): column
        for column in range(1, plan_sheet.max_column + 1)
    }
    missing_labels = [label for label in labels if label not in header_index]
    if missing_labels:
        raise ValueError(f"result2.xlsx缺少时间段：{missing_labels[:5]}")

    for row, day in enumerate(days, start=2):
        day_date = datetime.fromisoformat(day["date"])
        plan_sheet.cell(row, 1, day_date).number_format = "yyyy/m/d"
        purchases = day["planned_purchase_kwh"]
        if len(purchases) != PERIODS:
            raise ValueError(f"{day['date']}全天计划购电数组长度不为144")
        for label, value in zip(labels, purchases):
            cell = plan_sheet.cell(row, header_index[label], float(value))
            cell.number_format = "0.0000"
        plan_sheet.cell(row, 146, float(sum(purchases))).number_format = "0.0000"
        plan_sheet.cell(row, 147, float(day["planned_cost_yuan"])).number_format = "0.0000"

    battery_styles = [
        [copy(battery_sheet.cell(row, column)._style) for column in range(1, 7)]
        for row in range(2, 8)
    ]
    for row in battery_sheet.iter_rows(min_row=2, max_row=battery_sheet.max_row):
        for cell in row:
            cell.value = None
    for day_index, day in enumerate(days):
        charges = day["charge_kwh"]
        discharges = day["discharge_kwh"]
        states = day["soc_kwh"]
        if len(charges) != PERIODS or len(discharges) != PERIODS or len(states) != PERIODS + 1:
            raise ValueError(f"{day['date']}充放电或储能状态数组长度错误")
        for block in range(FOUR_HOUR_BLOCKS):
            row = 2 + day_index * FOUR_HOUR_BLOCKS + block
            for column, style in enumerate(battery_styles[block], start=1):
                battery_sheet.cell(row, column)._style = copy(style)
            battery_sheet.cell(row, 1).value = datetime.fromisoformat(day["date"]) if block == 0 else None
            battery_sheet.cell(row, 2).value = FOUR_HOUR_LABELS[block]
            battery_sheet.cell(row, 3, _sum_block(charges, block)).number_format = "0.0000"
            battery_sheet.cell(row, 4, _sum_block(discharges, block)).number_format = "0.0000"
            battery_sheet.cell(row, 5).value = "0:00" if block == 0 else "24:00" if block == 1 else None
            battery_sheet.cell(row, 6).value = float(states[0]) if block == 0 else float(states[-1]) if block == 1 else None
            battery_sheet.cell(row, 1).number_format = "yyyy/m/d"
            battery_sheet.cell(row, 6).number_format = "0.0000"
    if battery_sheet.max_row > 1 + len(days) * FOUR_HOUR_BLOCKS:
        battery_sheet.delete_rows(
            2 + len(days) * FOUR_HOUR_BLOCKS,
            battery_sheet.max_row - (1 + len(days) * FOUR_HOUR_BLOCKS),
        )

    emergency_rows = _build_emergency_rows(days)

    emergency_styles = [
        [copy(emergency_sheet.cell(row, column)._style) for column in range(1, 4)]
        for row in range(2, 5)
    ]
    for row in emergency_sheet.iter_rows(min_row=2, max_row=emergency_sheet.max_row):
        for cell in row:
            cell.value = None
    for output_index, (day_date, period, amount, group_index, group_size) in enumerate(emergency_rows):
        row = output_index + 2
        style_index = 0 if group_index == 0 else 2 if group_index == group_size - 1 else 1
        for column, style in enumerate(emergency_styles[style_index], start=1):
            emergency_sheet.cell(row, column)._style = copy(style)
        emergency_sheet.cell(row, 1, day_date).number_format = "yyyy/m/d"
        emergency_sheet.cell(row, 2).value = period
        emergency_sheet.cell(row, 3).value = amount
        emergency_sheet.cell(row, 3).number_format = "0.0000"
    if emergency_sheet.max_row > 1 + len(emergency_rows):
        emergency_sheet.delete_rows(
            2 + len(emergency_rows),
            emergency_sheet.max_row - (1 + len(emergency_rows)),
        )

    output = c_root / "结果" / "result2.xlsx"
    workbook.save(output)
    workbook.close()
    return output


def verify_problem1(c_root: Path) -> None:
    q1_solution = _read_json(c_root / "结果" / "q1_solution.json")
    workbook = load_workbook(c_root / "结果" / "result1.xlsx", data_only=False)
    plan_sheet = workbook["计划购电量"]
    battery_sheet = workbook["充放电量"]
    written_q1 = [float(plan_sheet.cell(row, 2).value) for row in range(2, PERIODS + 2)]
    if max(abs(a - b) for a, b in zip(written_q1, q1_solution["purchase_kwh"])) > 1e-9:
        raise AssertionError("result1.xlsx购电量写入校验失败")
    if abs(sum(written_q1) - q1_solution["total_purchase_kwh"]) > 1e-7:
        raise AssertionError("result1.xlsx购电量合计校验失败")
    if abs(float(battery_sheet["E2"].value) - 6000.0) > 1e-9 or abs(float(battery_sheet["E3"].value) - 6000.0) > 1e-9:
        raise AssertionError("result1.xlsx首末储能校验失败")
    workbook.close()


def verify_problem2(c_root: Path) -> None:
    q2_solution = _read_json(c_root / "结果" / "q2_solution.json")
    reporting_tolerance = float(
        q2_solution.get("model", {}).get("emergency_reporting_tolerance_kwh", 0.0)
    )
    if reporting_tolerance <= 0.0:
        raise AssertionError("问题二结果缺少有效的紧急购电报告阈值")
    workbook = load_workbook(c_root / "结果" / "result2.xlsx", data_only=False)
    plan_sheet = workbook["计划购电量"]
    battery_sheet = workbook["充放电量"]
    emergency_sheet = workbook["紧急购电量"]
    if plan_sheet.max_row != 335 or plan_sheet.max_column != 147:
        raise AssertionError("result2.xlsx全天计划表尺寸错误")
    header_index = {
        _normalize_result2_label(plan_sheet.cell(1, column).value): column
        for column in range(1, plan_sheet.max_column + 1)
    }
    for day_index, day in enumerate(q2_solution["days"]):
        if any(
            len(day[key]) != PERIODS
            for key in (
                "planned_purchase_kwh",
                "charge_kwh",
                "discharge_kwh",
                "emergency_kwh",
                "emergency_kwh_raw",
            )
        ) or len(day["soc_kwh"]) != PERIODS + 1:
            raise AssertionError(f"{day['date']}的逐时结果数组长度错误")
        if float(day["max_balance_error_kwh"]) > 1e-7:
            raise AssertionError(f"{day['date']}的原始LP供需平衡误差超限")
        if min(day["soc_kwh"]) < 1200.0 - 1e-7 or max(day["soc_kwh"]) > 10800.0 + 1e-7:
            raise AssertionError(f"{day['date']}的储能状态越界")
        if max(day["charge_kwh"] + day["discharge_kwh"]) > 5000.0 / 6.0 + 1e-7:
            raise AssertionError(f"{day['date']}的充放电功率越界")
        for reported, raw in zip(day["emergency_kwh"], day["emergency_kwh_raw"]):
            expected = float(raw) if float(raw) >= reporting_tolerance else 0.0
            if abs(float(reported) - expected) > 1e-12:
                raise AssertionError(f"{day['date']}的紧急购电报告阈值未统一")
        interval_total = sum(
            float(item["purchase_kwh"]) for item in day["emergency_intervals"]
        )
        if abs(interval_total - sum(day["emergency_kwh"])) > 1e-7:
            raise AssertionError(f"{day['date']}的紧急购电区间合计错误")
        plan_row = day_index + 2
        workbook_date = plan_sheet.cell(plan_row, 1).value
        if (
            not isinstance(workbook_date, datetime)
            or workbook_date.date().isoformat() != day["date"]
        ):
            raise AssertionError(f"{day['date']}的计划购电日期写入错误")
        for label, expected in zip(
            q2_solution["interval_labels"], day["planned_purchase_kwh"]
        ):
            actual = float(plan_sheet.cell(plan_row, header_index[label]).value)
            if abs(actual - float(expected)) > 1e-9:
                raise AssertionError(f"{day['date']}的{label}计划购电量写入错误")
        if (
            abs(
                float(plan_sheet.cell(plan_row, 146).value)
                - sum(day["planned_purchase_kwh"])
            )
            > 1e-7
        ):
            raise AssertionError(f"{day['date']}的全天计划购电量合计错误")
        if (
            abs(
                float(plan_sheet.cell(plan_row, 147).value)
                - float(day["planned_cost_yuan"])
            )
            > 1e-7
        ):
            raise AssertionError(f"{day['date']}的计划购电费写入错误")
        battery_row = 2 + day_index * FOUR_HOUR_BLOCKS
        if abs(float(battery_sheet.cell(battery_row, 6).value) - float(day["soc_kwh"][0])) > 1e-9:
            raise AssertionError(f"{day['date']}的0:00储能量写入错误")
        if abs(float(battery_sheet.cell(battery_row + 1, 6).value) - float(day["soc_kwh"][-1])) > 1e-7:
            raise AssertionError(f"{day['date']}的24:00储能量写入错误")
        if any(battery_sheet.cell(battery_row + offset, 1).value is not None for offset in range(1, FOUR_HOUR_BLOCKS)):
            raise AssertionError(f"{day['date']}的日期列存在模板残留值")
        if any(battery_sheet.cell(battery_row + offset, 5).value is not None for offset in range(2, FOUR_HOUR_BLOCKS)):
            raise AssertionError(f"{day['date']}的时刻列存在模板残留值")
        for block in range(FOUR_HOUR_BLOCKS):
            written_charge = float(battery_sheet.cell(battery_row + block, 3).value)
            written_discharge = float(battery_sheet.cell(battery_row + block, 4).value)
            if abs(written_charge - _sum_block(day["charge_kwh"], block)) > 1e-7:
                raise AssertionError(f"{day['date']}的第{block + 1}段充电量写入错误")
            if abs(written_discharge - _sum_block(day["discharge_kwh"], block)) > 1e-7:
                raise AssertionError(f"{day['date']}的第{block + 1}段放电量写入错误")
        if day_index > 0:
            previous_day = q2_solution["days"][day_index - 1]
            if abs(float(day["soc_kwh"][0]) - float(previous_day["soc_kwh"][-1])) > 1e-7:
                raise AssertionError(f"{day['date']}与前一日储能状态未连续传递")
    emergency_total = sum(
        float(emergency_sheet.cell(row, 3).value or 0.0)
        for row in range(2, emergency_sheet.max_row + 1)
    )
    if abs(emergency_total - q2_solution["summary"]["emergency_energy_kwh"]) > 1e-6:
        raise AssertionError("result2.xlsx紧急购电量合计校验失败")
    expected_emergency_rows = _build_emergency_rows(q2_solution["days"])
    if emergency_sheet.max_row != len(expected_emergency_rows) + 1:
        raise AssertionError("result2.xlsx紧急购电表行数错误")
    for row, (expected_date, expected_period, expected_amount, _, _) in enumerate(
        expected_emergency_rows, start=2
    ):
        actual_date = emergency_sheet.cell(row, 1).value
        actual_period = emergency_sheet.cell(row, 2).value
        actual_amount = emergency_sheet.cell(row, 3).value
        if actual_date != expected_date or actual_period != expected_period:
            raise AssertionError(f"result2.xlsx紧急购电表第{row}行日期或时段错误")
        if expected_amount is None:
            if actual_amount is not None:
                raise AssertionError(f"result2.xlsx紧急购电表第{row}行应为空")
        elif abs(float(actual_amount) - expected_amount) > 1e-7:
            raise AssertionError(f"result2.xlsx紧急购电表第{row}行购电量错误")
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith("#"):
                    raise AssertionError(f"{sheet.title}!{cell.coordinate}存在公式错误：{cell.value}")
                if cell.value == "⁝":
                    raise AssertionError(f"{sheet.title}!{cell.coordinate}存在模板省略符残留")
    workbook.close()


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="导出C题官方结果模板")
    parser.add_argument("question", choices=("1", "2", "all"), nargs="?", default="all")
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    c_root = Path(__file__).resolve().parents[1]
    outputs = []
    if args.question in ("1", "all"):
        outputs.append(export_problem1(c_root))
        verify_problem1(c_root)
    if args.question in ("2", "all"):
        outputs.append(export_problem2(c_root))
        verify_problem2(c_root)
    for output in outputs:
        print(f"已生成并校验：{output}")


if __name__ == "__main__":
    main()
